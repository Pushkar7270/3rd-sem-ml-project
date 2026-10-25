#!/usr/bin/env python3
"""Generate a viva pitch, Q&A playbook and glossary for the churn project.

Output:
    ~/Downloads/RetainIQ_Viva_Pitch_QA_Glossary.docx
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

OUTPUT = Path.home() / "Downloads" / "RetainIQ_Viva_Pitch_QA_Glossary.docx"

DARK = RGBColor(0x09, 0x1A, 0x14)
GREEN = RGBColor(0x0B, 0x9E, 0x68)
MUTED = RGBColor(0x5C, 0x73, 0x69)
RED = RGBColor(0xB3, 0x35, 0x45)
HEADER_FILL = "0B3B2E"
ZEBRA = "F4FAF7"
GREEN_TINT = "EBF9F2"
AMBER_TINT = "FFF7E8"


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def set_table_borders(table, color="D7E5DE", size=4):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)
    tbl_pr.append(borders)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_table(doc, headers, rows, widths=None, font_size=8.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    repeat_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, HEADER_FILL)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(str(value))
        run.bold = True
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.name = "Calibri"
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, ZEBRA)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            run = paragraph.add_run(str(value))
            run.font.size = Pt(font_size)
            run.font.name = "Calibri"
            run.font.color.rgb = DARK
    if widths:
        for index, width in enumerate(widths):
            for row in table.rows:
                row.cells[index].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_body(doc, text, size=10.4, italic=False, color=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.3
    run = paragraph.add_run(text)
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.2
        paragraph.add_run(item)


def add_steps(doc, steps):
    for index, step in enumerate(steps, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.55)
        paragraph.paragraph_format.first_line_indent = Cm(-0.45)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.2
        number = paragraph.add_run(f"{index}.")
        number.bold = True
        number.font.color.rgb = GREEN
        paragraph.add_run(f"   {step}")


def add_box(doc, title, text, fill=GREEN_TINT, border="BFE8D5", title_color=None):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, border, 6)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 130, 130, 170, 170)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = title_color or GREEN
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.22
    run = paragraph.add_run(text)
    run.font.size = Pt(9.4)
    run.font.color.rgb = DARK
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    heading.paragraph_format.space_before = Pt(13 if level == 1 else 8)
    heading.paragraph_format.space_after = Pt(5)
    heading.paragraph_format.keep_with_next = True
    return heading


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def style_document(doc):
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.4)
    normal.font.color.rgb = DARK
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(6)

    for level, size in ((1, 18), (2, 13), (3, 11)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.bold = True
        style.font.size = Pt(size)
        style.font.color.rgb = DARK if level == 1 else GREEN
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "RetainIQ  ·  Viva Pitch & Q&A Playbook"
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED


def cover(doc):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(70)
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, HEADER_FILL)
    set_cell_margins(cell, 190, 190, 220, 220)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("VIVA PREPARATION  ·  TOMORROW'S EVALUATION")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0xA7, 0xF7, 0xD5)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(36)
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run("RetainIQ")
    run.bold = True
    run.font.size = Pt(42)
    run.font.name = "Calibri"
    run.font.color.rgb = DARK

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(22)
    run = subtitle.add_run("Pitch, demo script, professor Q&A and glossary")
    run.font.size = Pt(17)
    run.font.color.rgb = GREEN
    run.bold = True

    add_body(
        doc,
        "A practical speaking guide for the Telco Customer Churn Intelligence project: what to say, "
        "how to demonstrate it, how to answer the likely questions, and the terminology you should "
        "know cold.",
        size=12,
        color=MUTED,
    )

    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(18)
    p_pr = rule._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2FF0A3")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    info = [
        ("Project", "Explainable churn prediction and retention recommendation system"),
        ("Dataset", "IBM Telco Customer Churn · 7,043 customers"),
        ("Model", "Logistic Regression · CV ROC-AUC 0.8475 · test ROC-AUC 0.8456"),
        ("Deployed system", "FastAPI + React 19 + SQLite + SHAP"),
        ("Prepared", f"{date.today():%d %B %Y}"),
    ]
    for label, value in info:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(f"{label}: ")
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = MUTED
        run = paragraph.add_run(value)
        run.font.size = Pt(9.5)
        run.font.color.rgb = DARK


def contents(doc):
    add_heading(doc, "Contents", 1)
    items = [
        ("1", "How to use this playbook"),
        ("2", "The 60-second answer"),
        ("3", "The six-minute pitch (word-for-word)"),
        ("4", "The five-minute live demo"),
        ("5", "Slide-by-slide talking points"),
        ("6", "Professor questions and strong answers"),
        ("7", "Curveball questions"),
        ("8", "Glossary"),
        ("9", "Numbers cheat sheet"),
        ("10", "Last-minute checklist"),
    ]
    for number, title in items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(f"{number}.")
        run.bold = True
        run.font.color.rgb = GREEN
        run = paragraph.add_run(f"   {title}")
        run.font.color.rgb = DARK


def section_howto(doc):
    add_heading(doc, "1. How to use this playbook", 1)
    add_body(
        doc,
        "Do not memorise every sentence. Learn the structure, rehearse the 60-second version until it "
        "is automatic, and know the five numbers in the cheat sheet. The pitch below is written in "
        "spoken English so it sounds natural when you say it out loud.",
    )
    add_box(
        doc,
        "Your core narrative in one line",
        "Prediction is not the product. The product is the path from a churn probability to an "
        "explanation, a risk band, a concrete retention action and a measurable outcome.",
    )
    add_bullets(
        doc,
        [
            "Open with the business problem, not with the algorithm.",
            "Use numbers to support claims: 7,043 customers, 26.5% churn, 0.8475 CV ROC-AUC, 74.1% recall.",
            "Always connect a technical choice to a business consequence.",
            "Admit limitations before the professor finds them; it signals maturity.",
            "If you do not know an answer, say so and describe how you would investigate it.",
        ],
    )
    add_heading(doc, "The five numbers to know cold", 2)
    add_table(
        doc,
        ["Number", "Meaning"],
        [
            ("7,043 / 26.5%", "Customers in the dataset / historical churn rate"),
            ("5,634 / 1,409", "Training rows / held-out test rows"),
            ("0.8475 / 0.8456", "5-fold CV ROC-AUC / test ROC-AUC for the selected model"),
            ("0.32 / 74.1%", "Deployed decision threshold / test recall at that threshold"),
            ("1,002 / 14.2%", "High-risk customers / share of the book"),
        ],
        widths=[1.8, 4.7],
    )


def section_short(doc):
    add_heading(doc, "2. The 60-second answer", 1)
    add_body(
        doc,
        "Use this when the professor says: “Tell me about your project” or “Give me the summary.” "
        "Speak slowly and land the last sentence.",
    )
    add_box(
        doc,
        "Say this",
        "Good morning, sir. My project is RetainIQ, an explainable customer-churn system built on the "
        "IBM Telco dataset. The problem is not only predicting who will leave; it is helping a retention "
        "team decide what to do next. I compared six classifiers under the same preprocessing and "
        "five-fold cross-validation. Logistic Regression won with a CV ROC-AUC of 0.8475 and achieved "
        "0.8456 on the held-out test set. At a threshold of 0.32 it catches about 74% of churners. Every "
        "score comes with signed SHAP drivers, a risk band, the revenue at risk and a personalised "
        "retention playbook. The full system scores all 7,043 customers, flags 1,002 as high risk, and "
        "exposes a React dashboard, a FastAPI backend and SQLite storage for actions and feedback. The "
        "main contribution is the complete path from prediction to explanation to action.",
    )
    add_body(doc, "If you only have 20 seconds:", italic=True, color=MUTED)
    add_box(
        doc,
        "Say this",
        "I built an explainable churn system on 7,043 Telco customers. Six models were compared with "
        "5-fold CV; the selected model reaches 0.8456 test ROC-AUC and 74% recall. It explains every score "
        "with SHAP and converts it into a retention action, not just a label.",
        fill=AMBER_TINT,
        border="F0D9A8",
        title_color=RGBColor(0x8A, 0x5A, 0x0A),
    )


def section_pitch(doc):
    add_heading(doc, "3. The six-minute pitch (word-for-word)", 1)
    add_body(
        doc,
        "This is a complete speaking script. Pause after each bold time marker. If you run long, skip "
        "the EDA detail and the model-family explanations; never skip the problem, the result or the "
        "business value.",
    )

    segments = [
        (
            "0:00–0:40 · Opening",
            "Good morning, sir. Thank you for the time. I will spend six minutes on the problem, the "
            "model, the evidence and the working system, and then I would be glad to take questions. "
            "My project is RetainIQ: an explainable customer-churn system. The one-line version is that "
            "I do not just predict churn; I explain it and turn it into a retention action.",
            "Do not start with “I used Logistic Regression”. Start with the problem.",
        ),
        (
            "0:40–1:30 · Problem and objectives",
            "Customer churn is the loss of an existing customer. It directly reduces recurring revenue "
            "and makes growth expensive because replacing a customer usually costs more than retaining "
            "one. Traditional churn analysis is slow, hard to reproduce and rarely explains why a "
            "particular customer is at risk. A basic model gives a probability but not a reason or an "
            "action. So I set four objectives: predict churn accurately, quantify the risk, explain the "
            "drivers, and recommend what the retention team should do.",
            "The gap is the last mile: probability → explanation → action.",
        ),
        (
            "1:30–2:20 · Data and EDA",
            "I used the IBM Telco Customer Churn dataset: 7,043 customers, 21 raw columns and a 26.5% "
            "churn rate. The EDA told a clear story. Month-to-month customers churn at about 42.7%, "
            "compared with 11.3% on one-year contracts and 2.8% on two-year contracts. Fiber-optic and "
            "electronic-check customers are also higher-risk segments. Churners have shorter tenure and "
            "higher monthly bills. These patterns shaped the retention playbook, not just the model.",
            "One concrete EDA number is more persuasive than a generic “I did EDA”.",
        ),
        (
            "2:20–3:10 · Method",
            "I cleaned the data and treated the 11 blank TotalCharges values as semantic zeros because "
            "they belong to customers with tenure zero. I engineered service-count features, a "
            "new-customer flag and a family flag. All preprocessing lives inside a scikit-learn Pipeline "
            "so imputation, scaling and one-hot encoding are fitted inside each cross-validation fold; "
            "that prevents data leakage. I compared Logistic Regression, Decision Tree, Random Forest, "
            "XGBoost, KNN and SVM with the same pipeline and stratified five-fold cross-validation, "
            "selecting on ROC-AUC.",
            "The key phrase is “same pipeline, same folds, same protocol”.",
        ),
        (
            "3:10–3:55 · Results",
            "Logistic Regression won with a mean CV ROC-AUC of 0.8475. XGBoost was close at 0.8409, "
            "Random Forest at 0.8359, Decision Tree 0.8265, KNN 0.8117 and SVM 0.7984. On the untouched "
            "test set the selected model reached 0.8456 ROC-AUC and 0.6519 PR-AUC. I tuned the threshold "
            "on out-of-fold predictions to maximise F1; the deployed threshold is 0.32. At that point "
            "recall is 74.1%, precision 54.1%, F1 0.625 and Brier score 0.136. The confusion matrix is "
            "800 true negatives, 235 false positives, 97 false negatives and 277 true positives.",
            "Explain why recall is high: missing a churner is more expensive than a wasted offer.",
        ),
        (
            "3:55–4:35 · Explainability",
            "A probability alone is not actionable, so I added SHAP. LinearSHAP is exact for the selected "
            "logistic model; TreeSHAP is used automatically for tree models and permutation SHAP is the "
            "fallback. Globally, the strongest drivers are internet service, contract, tenure and monthly "
            "charges. Locally, each customer gets signed contributions, and I show the category the "
            "customer actually holds. I also expose top_risk_driver: the strongest risk-increasing factor "
            "for customer-facing copy. I am careful to say that SHAP is conditional, not causal.",
            "This is the point where you distinguish yourself from a basic classifier.",
        ),
        (
            "4:35–5:15 · Business value and system",
            "The risk bands are aligned with the threshold: low below 0.32, medium from 0.32 to 0.60 and "
            "high at or above 0.60. In the current book, 4,552 customers are low risk, 1,489 medium and "
            "1,002 high. High-risk annual revenue at risk is about $705,663. A 14-rule playbook combines "
            "the risk band, account attributes and SHAP drivers to recommend actions such as annual "
            "contract conversion, fiber-plan review, autopay migration, onboarding or a targeted save "
            "offer. The system is FastAPI plus a React dashboard, with SQLite storing predictions, "
            "feedback and saved actions. It runs locally with no sign-in.",
            "Show that the model is embedded in a real workflow, not left in a notebook.",
        ),
        (
            "5:15–5:45 · Demo and impact",
            "If you open the dashboard, the overview shows the book, the risk mix, the model metrics and "
            "the priority queue. From a customer row I can show a 90% risk score, the tenure driver and "
            "the recommended action, then save that action and update its status in the queue. The score "
            "history records every prediction and lets me record whether the customer actually churned, "
            "which closes the feedback loop.",
            "Offer the demo, but do not let it run without narration.",
        ),
        (
            "5:45–6:00 · Limitations and close",
            "I would highlight three limitations. The data is observational, so SHAP is not causal. "
            "Probabilities are not post-hoc calibrated. The recommendations are heuristics, not measured "
            "uplift. The next steps are calibration, cost-sensitive thresholds, uplift testing and drift "
            "monitoring. Thank you, sir. I am happy to take questions.",
            "Ending with limitations makes you look confident, not unsure.",
        ),
    ]
    for title, say, note in segments:
        add_heading(doc, title, 2)
        add_box(doc, "Say this", say)
        add_body(doc, f"Delivery note: {note}", size=9.2, italic=True, color=MUTED)


def section_demo(doc):
    add_heading(doc, "4. The five-minute live demo", 1)
    add_body(
        doc,
        "Before you start: open http://127.0.0.1:8000/app/ in a clean browser tab, keep /docs open in a "
        "second tab, and have docs/screenshots/ open as a fallback. Reset the SQLite database if you "
        "want a clean queue, or leave one saved action so the queue is not empty.",
    )
    add_steps(
        doc,
        [
            "Overview (40 seconds): “This is the retention command center. The book has 7,043 customers, "
            "26.5% historical churn, 1,002 high-risk customers and about $705,663 of annual revenue at "
            "risk in the high band.”",
            "Risk distribution (20 seconds): point to the donut and say “the low/medium boundary is the "
            "deployed threshold, so a low-risk customer is never a predicted churner.”",
            "Customers (30 seconds): click Customers, choose High risk, and say “these are the customers "
            "the team should work first; the table ranks by probability and shows the top risk driver.”",
            "Open a customer (45 seconds): click the first row. “This customer has a 90% probability. "
            "The strongest risk-increasing driver is short tenure; monthly charges are protective for "
            "this profile. The model is conditional, not causal. Revenue at risk is $1,032 and the "
            "playbook starts with a fiber-plan review.”",
            "Save an action (30 seconds): click Add to queue, show the saved state, then open Retention "
            "queue. “The action is stored in SQLite; I can change its status to contacted or accepted.”",
            "Score a customer (60 seconds): open Score customer, click High-risk demo, click Score. “This "
            "is the same model on a live payload: 87.5% probability, high risk, with signed SHAP values "
            "and five recommended actions.” Point to the threshold and the top driver.",
            "Insights (45 seconds): open Model insights. “Six models, five-fold CV, selected on ROC-AUC. "
            "The deployed model is 0.8456 on test. Here is the ROC curve, the global SHAP drivers and the "
            "calibration/report gallery.”",
            "History (30 seconds): open Score history. “Every score is audited; I can record the actual "
            "outcome with Churned or Retained. That feedback is what would let us recalibrate and test "
            "retention offers later.”",
            "Close (20 seconds): “The system is a decision-support prototype. It predicts, explains, "
            "prioritises and recommends; it does not make autonomous offers.”",
        ],
    )
    add_box(
        doc,
        "If the demo fails",
        "Do not debug live. Say: “The API is restarting; I have the same screens captured in the "
        "documentation.” Open docs/screenshots/overview.png, customer_drawer.png and predict_result.png. "
        "Then continue the explanation. A calm fallback is better than five minutes of terminal errors.",
        fill=AMBER_TINT,
        border="F0D9A8",
        title_color=RGBColor(0x8A, 0x5A, 0x0A),
    )


def section_slides(doc):
    add_heading(doc, "5. Slide-by-slide talking points", 1)
    add_body(
        doc,
        "If you are using slides, keep each slide to one takeaway. These are the points to say, not the "
        "text to paste onto the slide.",
    )
    rows = [
        ("1. Title", "RetainIQ — explainable churn prediction and retention", "One sentence on the problem and the outcome."),
        ("2. Problem", "Churn costs revenue; basic models do not explain or act", "“The gap is the last mile.”"),
        ("3. Objectives", "Predict, explain, prioritise, recommend", "Four verbs, four deliverables."),
        ("4. Dataset", "7,043 · 21 columns · 26.5% churn", "Mention the semantic zero in TotalCharges."),
        ("5. EDA", "Contract, internet and payment patterns", "42.7% vs 2.8% for month-to-month vs two-year."),
        ("6. Method", "Leak-free pipeline + six models + 5-fold CV", "Same folds, same preprocessing, same protocol."),
        ("7. Results", "LR 0.8475 CV, 0.8456 test, 74% recall", "XGBoost was close; LR won on the pre-registered metric."),
        ("8. Explainability", "SHAP global + local, coefficients + importance", "Show one customer, one driver, one action."),
        ("9. Business value", "Risk bands, revenue at risk, 14-rule playbook", "$705,663 high-risk annual exposure."),
        ("10. System", "React + FastAPI + SQLite + tests", "Prediction → explanation → action → feedback."),
        ("11. Limitations", "Observational, uncalibrated, heuristic offers", "Say what you would do next."),
        ("12. Close", "Recap the contribution and invite questions", "Do not end on a limitation without a next step."),
    ]
    add_table(
        doc,
        ["Slide", "Say this", "Landing line"],
        rows,
        widths=[1.1, 2.6, 2.8],
        font_size=8.2,
    )


def add_question(doc, number, question, answer, push=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(f"Q{number}. {question}")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = DARK
    answer_paragraph = doc.add_paragraph()
    answer_paragraph.paragraph_format.left_indent = Cm(0.35)
    answer_paragraph.paragraph_format.space_after = Pt(3)
    answer_paragraph.paragraph_format.line_spacing = 1.22
    prefix = answer_paragraph.add_run("Answer: ")
    prefix.bold = True
    prefix.font.size = Pt(9.4)
    prefix.font.color.rgb = GREEN
    run = answer_paragraph.add_run(answer)
    run.font.size = Pt(9.4)
    run.font.color.rgb = RGBColor(0x28, 0x3B, 0x33)
    if push:
        push_paragraph = doc.add_paragraph()
        push_paragraph.paragraph_format.left_indent = Cm(0.35)
        push_paragraph.paragraph_format.space_after = Pt(3)
        run = push_paragraph.add_run(f"If he pushes: {push}")
        run.italic = True
        run.font.size = Pt(8.8)
        run.font.color.rgb = MUTED


def section_questions(doc):
    add_heading(doc, "6. Professor questions and strong answers", 1)
    add_body(
        doc,
        "These are the questions most likely to come from an ML evaluator. Answer in two to four "
        "sentences, then stop. If the professor wants more, he will ask a follow-up.",
    )

    add_heading(doc, "6.1 Fundamentals and framing", 2)
    add_question(doc, 1, "What is customer churn?",
                 "Churn is the loss of an existing customer through cancellation, non-renewal or migration "
                 "to a competitor. In this dataset it is the binary Churn column: Yes means the customer left.")
    add_question(doc, 2, "Why is churn prediction important for a business?",
                 "Recurring revenue falls and acquisition costs rise. Early identification lets the company "
                 "intervene before the customer leaves, and it lets the retention team prioritise limited "
                 "capacity instead of discounting everyone.")
    add_question(doc, 3, "Why use machine learning instead of fixed business rules?",
                 "Churn depends on many interacting attributes. A model learns multivariate and nonlinear "
                 "patterns and produces a probability. Rules are still useful for the action layer, which is "
                 "why my recommendations are rule-based and auditable.")
    add_question(doc, 4, "What is your contribution beyond training a classifier?",
                 "The end-to-end path: leak-free comparison of six models, threshold and risk-band alignment, "
                 "SHAP explanations, a 14-rule retention playbook, a FastAPI backend, a React dashboard, "
                 "SQLite audit/feedback and 166 tests. The classifier is only one component.")
    add_question(doc, 5, "Why did you choose this dataset?",
                 "It is public, structured, sufficiently large and contains the attributes the brief specifies: "
                 "tenure, contract, monthly charges, payment method, internet services and the Churn target. "
                 "It is also widely used, which makes the results comparable.")

    add_heading(doc, "6.2 Data and preprocessing", 2)
    add_question(doc, 6, "Describe the dataset.",
                 "IBM Telco Customer Churn: 7,043 customers, 21 raw columns, one row per customer and a 26.5% "
                 "churn rate. The target is binary; the features cover account, services, contract and billing.")
    add_question(doc, 7, "How did you handle the missing TotalCharges values?",
                 "There are 11 blanks, all for customers with tenure zero. Their lifetime spend is genuinely "
                 "zero, so I filled them with 0 rather than the median. Median imputation would invent revenue "
                 "that never existed.")
    add_question(doc, 8, "What did the EDA tell you?",
                 "Month-to-month customers churn at 42.7%, one-year at 11.3% and two-year at 2.8%. Fiber-optic "
                 "and electronic-check customers are also higher-risk. Churners have shorter tenure and higher "
                 "monthly bills. Those patterns shaped the playbook.")
    add_question(doc, 9, "Why not use customerID as a feature?",
                 "It is an arbitrary identifier with no generalisable signal. Using it would risk memorisation "
                 "and leakage, and it would not work for a new customer.")
    add_question(doc, 10, "Which features did you engineer and why?",
                  "NumServices, SupportServices and StreamingServices capture product depth and switching cost; "
                  "IsNewCustomer captures the high-risk first six months; HasFamily captures household stability.")
    add_question(doc, 11, "Why did you remove AvgMonthlyCharges and ChargeGap?",
                  "They are nearly collinear with tenure and monthly charges. They produced unstable coefficients "
                  "and counter-intuitive SHAP values for less than 0.001 ROC-AUC. Removing them made the model "
                  "simpler and the explanations defensible.")
    add_question(doc, 12, "How did you split the data?",
                  "A stratified 80/20 split with random_state 42: 5,634 training customers and 1,409 test "
                  "customers. Stratification preserves the 26.5% churn rate in both parts.")
    add_question(doc, 13, "What is data leakage and how did you prevent it?",
                  "Leakage is using information that would not be available at prediction time, or fitting "
                  "preprocessing on the test data. Every imputer, scaler and encoder lives inside a Pipeline that "
                  "is cloned and fitted inside each CV fold; the test set is touched once, after selection.")
    add_question(doc, 14, "Why use a Pipeline?",
                  "It bundles imputation, scaling, encoding and the classifier into one reproducible object. "
                  "It makes training and serving identical and makes leakage much harder.")
    add_question(doc, 15, "How do you handle categorical variables and unseen categories?",
                  "Most-frequent imputation plus one-hot encoding. At serving time, handle_unknown='ignore' "
                  "emits all zeros for an unseen category instead of crashing the API.")
    add_question(doc, 16, "Why scale the features?",
                  "KNN and SVM are distance-sensitive, so scaling is required. It is harmless for trees and keeps "
                  "one preprocessing path for a fair comparison.")

    add_heading(doc, "6.3 Models and training", 2)
    add_question(doc, 17, "Why these six models?",
                  "They cover the main families: linear, single tree, bagging, boosting, distance-based and "
                  "margin-based. That gives a fair comparison of bias/variance trade-offs on the same data.")
    add_question(doc, 18, "How does Logistic Regression work?",
                  "It models the log-odds of churn as a linear combination of features: log(p/(1-p)) = β0 + βx. "
                  "The sigmoid converts that to a probability, and coefficients are log-odds changes.")
    add_question(doc, 19, "Why did Logistic Regression beat XGBoost?",
                  "It had the highest mean CV ROC-AUC: 0.8475 versus 0.8409 for XGBoost. The dataset has strong "
                  "linear signal in contract, tenure and service type. XGBoost was close, but the extra "
                  "nonlinearity did not justify the complexity, and the linear model is easier to explain.")
    add_question(doc, 20, "Why not use a deep learning model?",
                  "This is small, tabular data with 7,043 rows. Deep networks need more data and tuning to justify "
                  "themselves, and they are harder to explain. The tree ensembles and linear model are the right "
                  "inductive bias here. A neural network could be future work, but it is not automatically better.")
    add_question(doc, 21, "Why were KNN and SVM weaker?",
                  "KNN suffers from distance concentration in a high-dimensional one-hot space. The RBF SVM is "
                  "competitive but slightly behind and harder to explain. Both are still useful baselines.")
    add_question(doc, 22, "Did you tune hyperparameters?",
                  "The comparison uses sensible, comparable defaults; the goal was a fair model-family comparison, "
                  "not an exhaustive search. The decision threshold was tuned. Nested cross-validation and "
                  "randomized search are the next step.")
    add_question(doc, 23, "What is 5-fold cross-validation?",
                  "The training set is split into five stratified folds. Each fold is held out once while the "
                  "model trains on the other four, and the five scores are averaged. It gives a more stable "
                  "estimate than a single validation split.")
    add_question(doc, 24, "Why five folds?",
                  "Five is a standard bias/variance compromise: each model sees 80% of the data per fit, and the "
                  "five estimates are enough to report a mean and standard deviation without excessive runtime.")
    add_question(doc, 25, "How did you handle class imbalance?",
                  "The positive class is 26.5%. I used stratified sampling, ROC-AUC and PR-AUC, and an F1-tuned "
                  "threshold instead of accuracy. I did not use SMOTE; class weights, SMOTE and cost-sensitive "
                  "learning are future comparisons.")
    add_question(doc, 26, "Why not optimise accuracy?",
                  "A model that predicts ‘stay’ for everyone is about 73% accurate and catches zero churners. "
                  "Accuracy hides the minority class; ROC-AUC, PR-AUC, recall and the confusion matrix are more "
                  "informative.")

    add_heading(doc, "6.4 Metrics and evaluation", 2)
    add_question(doc, 27, "What is ROC-AUC?",
                  "It is the probability that a randomly chosen churner is ranked above a randomly chosen "
                  "stayer. It measures ranking quality across all thresholds and is insensitive to the class "
                  "prior, which is why it was the primary selection metric.")
    add_question(doc, 28, "What is PR-AUC and why report it?",
                  "The area under the precision–recall curve. It focuses on the minority positive class and is "
                  "more informative than ROC-AUC when churn is only 26.5%. The test PR-AUC is 0.6519 against a "
                  "0.265 baseline.")
    add_question(doc, 29, "Precision or recall — which matters more?",
                  "Recall usually matters more here because a missed churner is more expensive than a wasted "
                  "retention offer. The deployed threshold gives 74.1% recall and 54.1% precision. The right "
                  "trade-off depends on the cost of each error.")
    add_question(doc, 30, "Why is your threshold 0.32?",
                  "I tuned it on out-of-fold training predictions to maximise F1. It raises recall from 52.7% at "
                  "0.5 to 74.1%, at the cost of some precision. The threshold is stored in the artifact and "
                  "returned by the API.")
    add_question(doc, 31, "What does the confusion matrix look like?",
                  "At the deployed threshold: 800 true negatives, 235 false positives, 97 false negatives and "
                  "277 true positives. The model catches 277 of 374 churners while accepting 235 false alarms.")
    add_question(doc, 32, "What is F1?",
                  "The harmonic mean of precision and recall. It balances the two error types and is useful when "
                  "both matter. The test F1 at the tuned threshold is 0.625.")
    add_question(doc, 33, "What is the Brier score?",
                  "The mean squared error of predicted probabilities against the 0/1 outcome. Lower is better; "
                  "0.136 means the probabilities are reasonably informative. It also rewards calibration, not "
                  "just ranking.")
    add_question(doc, 34, "Is the model overfitting?",
                  "The CV ROC-AUC is 0.8475 and the test ROC-AUC is 0.8456, a gap of only 0.0019. Every "
                  "preprocessing step is fitted inside the folds, and the test set was untouched until the end. "
                  "That is strong evidence against overfitting.")
    add_question(doc, 35, "Did you use the test set for model selection?",
                  "No. Selection used only 5-fold CV on the training set. The test set was evaluated once after "
                  "selection and threshold tuning, which keeps the final estimate honest.")
    add_question(doc, 36, "Is 0.85 ROC-AUC good enough?",
                  "It is competitive for this dataset and materially better than random. The bigger point is the "
                  "last mile: even a good probability is useless without explanation and action. I would still "
                  "improve calibration and test uplift before claiming business impact.")

    add_heading(doc, "6.5 Explainability", 2)
    add_question(doc, 37, "What is SHAP?",
                  "SHAP assigns each feature a contribution to a prediction using Shapley values from cooperative "
                  "game theory. Contributions are additive: base value plus the feature contributions reconstructs "
                  "the model output. LinearSHAP is exact for the selected logistic model.")
    add_question(doc, 38, "What is the difference between global and local explanations?",
                  "Global importance averages absolute SHAP values across customers and shows what drives the "
                  "model overall. Local explanation shows the signed drivers for one customer. Both are exposed: "
                  "global in Model insights, local in the customer drawer and /predict.")
    add_question(doc, 39, "Why can monthly charges have a negative SHAP value?",
                  "SHAP is conditional. For a new fiber customer with streaming add-ons, the rest of the profile "
                  "already explains the risk; the incremental price contribution can be protective. It is not a "
                  "marginal causal effect, and I say that explicitly in the UI.")
    add_question(doc, 40, "Coefficients or SHAP — which is better?",
                  "Coefficients are global and exact for the linear model, but they do not explain a single "
                  "prediction. SHAP is local and additive. I report both and use top_risk_driver for customer "
                  "copy so a protective factor is never presented as a reason to call.")
    add_question(doc, 41, "What are the top churn drivers?",
                  "Internet service, contract, tenure, monthly charges, streaming services, payment method, "
                  "new-customer status and online security. Contract and internet service are also the strongest "
                  "retention levers.")

    add_heading(doc, "6.6 Business, system and delivery", 2)
    add_question(doc, 42, "How are the risk bands defined?",
                  "Low is below 0.32, medium is 0.32–0.60 and high is 0.60 or above. The low/medium boundary is "
                  "the deployed threshold, so a customer is never both low risk and a predicted churner.")
    add_question(doc, 43, "How are the recommendations generated?",
                  "A 14-rule playbook combines the risk band, contract/service attributes and SHAP driver "
                  "alignment, then ranks actions by a priority score. Examples: annual-contract conversion, fiber "
                  "plan review, autopay migration, tech-support trial, onboarding and a targeted save offer.")
    add_question(doc, 44, "Are the recommendations causal?",
                  "No. They are auditable heuristics, not measured uplift. An expensive offer should be validated "
                  "with an A/B test or an uplift model. The system is decision support, not an autonomous agent.")
    add_question(doc, 45, "How do you calculate revenue at risk and CLV?",
                  "Annual revenue at risk = MonthlyCharges × 12 × P(churn). Expected remaining months = "
                  "min(1/P(churn), 72), and estimated CLV = MonthlyCharges × expected months. The CLV is a "
                  "revenue-only upper bound, not an accounting figure.")
    add_question(doc, 46, "Why FastAPI and SQLite?",
                  "FastAPI gives typed request/response contracts, automatic OpenAPI docs and a clean service "
                  "layer. SQLite is embedded in Python, needs no server, credentials or sign-in, and gives ACID "
                  "storage for the prediction audit trail, feedback and saved actions. Postgres is the production "
                  "swap.")
    add_question(doc, 47, "Why React and Vite?",
                  "React gives a component model for the dashboard, and Vite gives a fast dev server and a small "
                  "production bundle. The compiled bundle is served by FastAPI, so Node is not needed at demo "
                  "time. The UI uses glassmorphism but keeps contrast and a no-backdrop-filter fallback.")
    add_question(doc, 48, "How is the system tested?",
                  "166 automated tests cover data cleaning, features, all six models, metrics, SHAP, all 14 "
                  "recommendation rules, risk formulas, API contracts, history/feedback, actions and the static "
                  "frontend. A Playwright smoke test ran the full browser flow with zero console errors.")
    add_question(doc, 49, "What was the hardest part?",
                  "Aligning the explanation, the risk band and the action so they do not contradict each other. "
                  "The threshold had to match the low-risk boundary, SHAP drivers had to show the customer's actual "
                  "category, and recommendations had to suppress offers for services the customer does not have.")
    add_question(doc, 50, "How did the team divide the work?",
                  "Pushkar led ML and backend (45%), Karan owned research and frontend (35%), and Shaurya "
                  "supported data validation, testing and documentation (20%).")

    add_heading(doc, "6.7 Limitations, ethics and future work", 2)
    add_question(doc, 51, "What are the limitations?",
                  "Observational data, so no causal claims; probabilities are not post-hoc calibrated; the "
                  "threshold optimises F1 rather than monetary cost; recommendations are heuristics; the dataset "
                  "is US-centric; and the demo has no authentication by design.")
    add_question(doc, 52, "What ethical concerns exist?",
                  "Gender and senior status are present. The model should be audited for disparate impact and used "
                  "to target support, never to deny service or price discriminatorily. Customer-facing offers need "
                  "a human in the loop.")
    add_question(doc, 53, "What would you do with more time?",
                  "Probability calibration, cost-sensitive threshold optimisation per channel, uplift modelling, "
                  "nested hyperparameter search, drift monitoring, scheduled retraining and a production database "
                  "with auth and role-based access.")
    add_question(doc, 54, "How would you deploy this in production?",
                  "Containerise the API, put Postgres behind the same repository interface, add authentication "
                  "and RBAC, add a job queue for batch scoring, monitor input drift and calibration, and use a "
                  "model registry plus CI/CD for retraining and rollback.")


def section_curveballs(doc):
    add_heading(doc, "7. Curveball questions", 1)
    add_body(
        doc,
        "These are the questions that can unsettle a candidate. The goal is not a perfect answer; it is a "
        "calm, structured answer.",
    )
    curveballs = [
        (
            "“Your AUC is only 0.85. Why should we be impressed?”",
            "“That is a fair challenge. Churn is noisy and the dataset has a practical ceiling; 0.85 is "
            "competitive for Telco and the CV/test gap is only 0.002. But the differentiator is not the "
            "last decimal of AUC — it is that the system explains the score and converts it into an "
            "action. I would still improve calibration and test uplift before claiming business impact.”",
        ),
        (
            "“Isn’t this just a Kaggle notebook with a dashboard?”",
            "“The model is only one layer. I built a reproducible comparison protocol, an aligned "
            "threshold and risk policy, exact SHAP explanations, a 14-rule playbook, a FastAPI service, "
            "SQLite audit/feedback, a React UI and 166 tests. The contribution is the full decision-support "
            "workflow, not a notebook score.”",
        ),
        (
            "“Why should I trust SHAP?”",
            "“For the selected linear model, LinearSHAP is exact and additive; I verify additivity in the "
            "tests. For tree models I use TreeSHAP. I also cross-check with coefficients and permutation "
            "importance. I explicitly label SHAP as conditional, not causal, and use top_risk_driver for "
            "customer-facing copy.”",
            ),
        (
            "“If the model says a customer will churn, will they?”",
            "“No. It is a probability, not a prophecy. At the deployed threshold, 277 of 374 churners are "
            "caught, but there are 97 misses and 235 false alarms. The value is prioritisation, not "
            "certainty.”",
        ),
        (
            "“What happens when the data distribution changes?”",
            "“That is drift. I would monitor feature distributions, prediction distributions, calibration "
            "and the observed churn rate from the feedback table, then retrain on a schedule or when "
            "performance crosses a threshold. The current system already records predictions and outcomes "
            "to make that possible.”",
        ),
        (
            "“Why no authentication?”",
            "“Authentication is deliberately out of scope for the local viva demo; adding a login screen "
            "would not improve the ML. In production I would add OAuth/OIDC, role-based access, HTTPS, "
            "audit logging and tenant isolation before exposing customer data.”",
        ),
        (
            "“Can you prove the recommendations work?”",
            "“Not yet, and I would not claim that. They are a prioritised playbook based on churn drivers "
            "and business logic. Proving impact requires an A/B test or an uplift model; the feedback "
            "table is the first step toward that.”",
        ),
        (
            "“Why is there no probability calibration plot in the API?”",
            "“The training report includes a reliability diagram and Brier score; the API exposes the Brier "
            "score and raw probabilities. Post-hoc isotonic or Platt calibration is a listed next step, "
            "especially before using these probabilities in a financial decision.”",
        ),
    ]
    for index, (question, answer) in enumerate(curveballs, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(f"C{index}. {question}")
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RED
        answer_paragraph = doc.add_paragraph()
        answer_paragraph.paragraph_format.left_indent = Cm(0.4)
        answer_paragraph.paragraph_format.space_after = Pt(4)
        answer_paragraph.paragraph_format.line_spacing = 1.22
        run = answer_paragraph.add_run(answer)
        run.font.size = Pt(9.4)
        run.font.color.rgb = RGBColor(0x28, 0x3B, 0x33)


def section_glossary(doc):
    add_heading(doc, "8. Glossary", 1)
    add_body(
        doc,
        "If the professor uses a term you do not recognise, ask him to repeat it and define it in one "
        "sentence. These are the terms you should be able to explain without notes.",
    )
    terms = [
        ("Churn", "Loss of an existing customer through cancellation or non-renewal.", "The target event."),
        ("Churn rate", "Share of customers who churn in a period; here 26.5%.", "The base rate and business KPI."),
        ("Supervised learning", "Learning a mapping from labelled inputs to a target.", "The project has historical Yes/No labels."),
        ("Binary classification", "Predicting one of two classes.", "Churn vs stay."),
        ("Feature", "An input variable used by the model.", "Tenure, contract, monthly charges."),
        ("Target / label", "The value the model learns to predict.", "Churn = 1/0."),
        ("Class imbalance", "Unequal class frequencies.", "Only 26.5% churn, so accuracy is misleading."),
        ("Train/test split", "Separate data for fitting and final evaluation.", "80/20 here."),
        ("Stratified split", "A split preserving the class ratio.", "Keeps churn at 26.5% in both parts."),
        ("Data leakage", "Using information unavailable at prediction time, or fitting on test data.", "Prevented by pipelines inside CV."),
        ("Overfitting", "Learning noise and failing to generalise.", "CV/test gap is only 0.002."),
        ("Cross-validation", "Repeated train/validation splits to estimate generalisation.", "5-fold CV."),
        ("Stratified K-fold", "K-fold CV preserving class balance in each fold.", "Stable ROC-AUC estimates."),
        ("Pipeline", "A single object chaining preprocessing and the classifier.", "Prevents leakage and keeps serving identical."),
        ("Imputation", "Filling missing values.", "Median for numeric, most-frequent for categorical."),
        ("Standardisation", "Scaling to zero mean and unit variance.", "Required by KNN/SVM."),
        ("One-hot encoding", "Turning a categorical value into binary indicator columns.", "Contract, payment, services."),
        ("Hyperparameter", "A setting chosen before fitting.", "Tree depth, k, C, learning rate."),
        ("Logistic Regression", "Linear model of log-odds through a sigmoid.", "Selected model; interpretable coefficients."),
        ("Decision Tree", "Recursive rule-based partitioning.", "Interpretable but high variance."),
        ("Random Forest", "Bagged ensemble of trees.", "Strong tabular baseline."),
        ("Gradient Boosting / XGBoost", "Trees built sequentially to correct errors.", "Close second; more complex."),
        ("KNN", "Predicts from the nearest labelled neighbours.", "Sensitive to scaling and high dimensions."),
        ("SVM", "Margin classifier; RBF kernel here.", "Calibrated probabilities; slightly weaker."),
        ("ROC curve", "TPR vs FPR across thresholds.", "Shows ranking quality."),
        ("ROC-AUC", "Area under the ROC curve.", "Primary selection metric; 0.8475 CV."),
        ("PR curve", "Precision vs recall across thresholds.", "Better for minority classes."),
        ("PR-AUC", "Area under the PR curve.", "Test 0.6519 vs 0.265 baseline."),
        ("Precision", "Of predicted churners, the share that truly churn.", "54.1% at the deployed threshold."),
        ("Recall", "Of actual churners, the share the model catches.", "74.1% at the deployed threshold."),
        ("F1", "Harmonic mean of precision and recall.", "0.625 on test."),
        ("Accuracy", "Share of all predictions that are correct.", "76.4%, but misleading under imbalance."),
        ("Confusion matrix", "Counts of TP, FP, TN, FN.", "800/235/97/277 at the threshold."),
        ("Decision threshold", "Probability cut-off for class assignment.", "0.32, F1-tuned."),
        ("Calibration", "Agreement between predicted probabilities and observed frequencies.", "Brier 0.136; post-hoc calibration is future work."),
        ("Brier score", "Mean squared error of probabilities.", "Lower is better; rewards calibration."),
        ("Log loss", "Penalty for confident wrong probabilities.", "0.4156 on test."),
        ("SHAP", "Shapley-value attribution of a prediction to features.", "Additive local explanation."),
        ("LinearSHAP / TreeSHAP", "Exact SHAP variants for linear/tree models.", "Used automatically by model family."),
        ("Feature importance", "Native tree measure of split usefulness.", "Cross-check for SHAP."),
        ("Coefficient", "Log-odds weight in a linear model.", "Global and signed for LR."),
        ("Risk band", "Business grouping of probabilities.", "Low <0.32, medium 0.32–0.60, high ≥0.60."),
        ("CLV", "Customer lifetime value; expected revenue over time.", "Revenue-only upper bound here."),
        ("Revenue at risk", "Monthly charges × 12 × P(churn).", "$705,663 high-risk annual exposure."),
        ("Uplift", "Causal effect of an intervention on retention.", "Not measured; future A/B work."),
        ("A/B test", "Randomised comparison of two treatments.", "How offers should be validated."),
        ("FastAPI", "Python web framework for the API.", "Typed endpoints + OpenAPI docs."),
        ("Pydantic", "Data validation library used by FastAPI.", "Request/response contracts."),
        ("REST API", "HTTP interface for prediction and data.", "/predict, /customers, /actions."),
        ("SQLite", "Embedded, serverless SQL database.", "No install or sign-in; stores audit/actions."),
        ("React", "Component-based JavaScript UI library.", "The dashboard frontend."),
        ("Vite", "Frontend build tool and dev server.", "Builds the production bundle."),
        ("SPA", "Single-page application.", "The dashboard does not reload the page."),
        ("CORS", "Browser policy for cross-origin requests.", "Open for local dev; restrict in production."),
        ("WAL", "SQLite write-ahead logging.", "Better concurrency for reads/writes."),
        ("Audit trail", "Chronological record of predictions and decisions.", "The predictions table."),
        ("Glassmorphism", "Frosted-glass visual style using translucency and blur.", "Used in the UI; readability preserved."),
        ("Data drift", "Change in input distribution over time.", "Monitor and retrain."),
    ]
    add_table(
        doc,
        ["Term", "Meaning", "Why it matters"],
        terms,
        widths=[1.5, 3.0, 2.0],
        font_size=7.4,
    )


def section_numbers(doc):
    add_heading(doc, "9. Numbers cheat sheet", 1)
    add_body(doc, "Glance at this immediately before you walk in.")
    rows = [
        ("Dataset", "IBM Telco Customer Churn"),
        ("Rows / raw columns", "7,043 / 21"),
        ("Churn rate", "26.5% (1,869 churners)"),
        ("Train / test", "5,634 / 1,409 (stratified 80/20)"),
        ("Models compared", "Logistic Regression, Decision Tree, Random Forest, XGBoost, KNN, SVM"),
        ("CV ROC-AUC", "LR 0.8475 ± 0.0118; XGBoost 0.8409; RF 0.8359; DT 0.8265; KNN 0.8117; SVM 0.7984"),
        ("Selected model", "Logistic Regression (5-fold CV ROC-AUC)"),
        ("Test ROC-AUC / PR-AUC", "0.8456 / 0.6519"),
        ("Threshold", "0.32 (F1-tuned on out-of-fold predictions)"),
        ("Recall / precision / F1", "0.7406 / 0.5410 / 0.6253"),
        ("Accuracy / Brier / log loss", "0.7644 / 0.1362 / 0.4156"),
        ("Confusion matrix", "TN 800 · FP 235 · FN 97 · TP 277"),
        ("Risk distribution", "4,552 low · 1,489 medium · 1,002 high"),
        ("High-risk revenue at risk", "$705,663 per year"),
        ("Top drivers", "Internet service, contract, tenure, monthly charges, streaming, payment method"),
        ("Recommendation rules", "14 rules, 5 returned by default"),
        ("System", "FastAPI + React 19 + Vite + SQLite + SHAP"),
        ("Tests", "166 automated tests, Ruff clean, Playwright UI smoke test"),
        ("Team contributions", "Pushkar 45% · Karan 35% · Shaurya 20%"),
    ]
    add_table(doc, ["Item", "Value"], rows, widths=[1.8, 4.7], font_size=8.2)


def section_checklist(doc):
    add_heading(doc, "10. Last-minute checklist", 1)
    add_heading(doc, "15 minutes before", 2)
    add_steps(
        doc,
        [
            "Start the API: `cd ~/myWork/3rd-sem-ml-project && uv run churn-serve`.",
            "Open http://127.0.0.1:8000/app/ and confirm the sidebar says ‘Model online’.",
            "Open http://127.0.0.1:8000/docs in a second tab.",
            "Open docs/screenshots/ in the file manager as a demo fallback.",
            "Set browser zoom to 100% and close unrelated tabs.",
            "If you want a clean queue, reset data/processed/churn.db before starting; otherwise leave one saved action.",
            "Keep the terminal visible but not in the presentation window.",
        ],
    )
    add_heading(doc, "What to carry", 2)
    add_bullets(
        doc,
        [
            "This playbook, with the numbers page open.",
            "The Word project documentation and, ideally, a PDF copy.",
            "A USB copy of the repository, the trained artifact and the screenshots.",
            "A charger and a browser with the local URL bookmarked.",
        ],
    )
    add_heading(doc, "If you get stuck", 2)
    add_bullets(
        doc,
        [
            "If you do not know a term: ask the professor to repeat it, then give the definition you do know.",
            "If a number is challenged: point to the model card or metadata JSON.",
            "If the demo fails: switch to screenshots and keep narrating; do not debug live.",
            "If a question is outside scope: say so, then explain how you would investigate it.",
            "Never bluff a metric. Say ‘I don’t remember the exact value, but the direction is…’.",
        ],
    )
    add_heading(doc, "Three questions you can ask the professor", 2)
    add_bullets(
        doc,
        [
            "“Would you prioritise probability calibration or uplift modelling as the next step?”",
            "“Which cost assumption should drive the threshold in a real retention team — missed churners or wasted offers?”",
            "“Would you recommend exposing SHAP directly to agents, or only the top risk driver?”",
        ],
    )
    add_box(
        doc,
        "Final reminder",
        "The professor is evaluating whether you understand the problem and the trade-offs, not whether "
        "you memorised every library parameter. Lead with the business problem, support it with numbers, "
        "be honest about limitations, and finish with the working system.",
    )


def build():
    doc = Document()
    style_document(doc)
    cover(doc)
    doc.add_page_break()
    contents(doc)
    doc.add_page_break()
    section_howto(doc)
    doc.add_page_break()
    section_short(doc)
    doc.add_page_break()
    section_pitch(doc)
    doc.add_page_break()
    section_demo(doc)
    doc.add_page_break()
    section_slides(doc)
    doc.add_page_break()
    section_questions(doc)
    doc.add_page_break()
    section_curveballs(doc)
    doc.add_page_break()
    section_glossary(doc)
    doc.add_page_break()
    section_numbers(doc)
    doc.add_page_break()
    section_checklist(doc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Saved {OUTPUT}")


if __name__ == "__main__":
    build()
